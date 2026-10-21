# -*- coding: utf-8 -*-
"""docx_inspector.core — 数据模型、规则注册表、文档上下文

设计约束：
  * 只做确定性判断，不涉及审美或内容评价
  * 每条规则独立、可单独开关、可单独测试
  * 输出 finding 必须带机器可读的 location 与 detail
"""
from __future__ import annotations

import os
import zipfile
from dataclasses import dataclass, field, asdict
from typing import Callable, Iterable

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu

SEVERITY_ORDER = {"error": 0, "warning": 1, "info": 2}


# ------------------------------------------------------------------ 模型
@dataclass
class Finding:
    rule: str
    severity: str          # error | warning | info
    message: str
    location: str = ""     # 人类可读定位，如 "第 3 个表格" / "第 1 页"
    fixable: bool = False
    detail: dict = field(default_factory=dict)

    def to_dict(self):
        d = asdict(self)
        return d

    def to_line(self):
        tag = {"error": "ERROR", "warning": "WARN", "info": "INFO"}[self.severity]
        loc = ("  " + self.location) if self.location else ""
        return "%-6s %-5s %s%s" % (tag, self.rule, self.message, loc)


@dataclass
class RuleMeta:
    id: str
    system: str            # package | geometry | font | layout | refs
    title: str
    severity: str
    default_on: bool
    fixable: bool
    fn: Callable = None


RULES: dict[str, RuleMeta] = {}


def rule(rid: str, system: str, title: str, severity: str = "warning",
         default_on: bool = True, fixable: bool = False):
    """注册一条规则。被装饰函数签名 fn(ctx) -> list[Finding]"""
    def deco(fn):
        RULES[rid] = RuleMeta(rid, system, title, severity, default_on, fixable, fn)
        return fn
    return deco


# ------------------------------------------------------------------ 上下文
class Ctx:
    """一次性算出各规则都要用的公共信息，避免重复打开/解析"""

    def __init__(self, path: str):
        self.path = os.path.abspath(path)
        self.size = os.path.getsize(self.path)

        # ZIP 层（不解压，直接看条目）
        self.zf = zipfile.ZipFile(self.path)
        self.names = self.zf.namelist()
        self.name_set = set(self.names)
        self._bad_crc = None

        # 文档层
        self.doc = Document(self.path)
        doc = self.doc
        sec = doc.sections[0]
        self.section = sec
        self.page_w = Emu(sec.page_width)
        self.page_h = Emu(sec.page_height)
        self.margin_l = Emu(sec.left_margin)
        self.margin_r = Emu(sec.right_margin)
        self.margin_t = Emu(sec.top_margin)
        self.margin_b = Emu(sec.bottom_margin)
        self.usable_w = Emu(int(sec.page_width - sec.left_margin - sec.right_margin))
        self.usable_h = Emu(int(sec.page_height - sec.top_margin - sec.bottom_margin))

        self._fonts_used = None
        self._used_styles = None
        self._installed = None

    # -- 惰性计算 -------------------------------------------------
    @property
    def used_style_names(self) -> set:
        """文档实际引用到的样式名集合。

        模板里常带一堆未使用的内置样式（macro / HTML Preformatted 等），
        它们引用的字体和 theme 属性与实际排版无关，检查它们只会制造误报。
        """
        if self._used_styles is not None:
            return self._used_styles
        doc = self.doc
        names = {p.style.name for p in doc.paragraphs if p.style is not None}
        for p in doc.paragraphs:
            for r in p.runs:
                if r.style is not None:
                    names.add(r.style.name)
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.style is not None:
                            names.add(p.style.name)
                        for r in p.runs:
                            if r.style is not None:
                                names.add(r.style.name)
        for sec in doc.sections:
            for hf in (sec.header, sec.footer, sec.first_page_header,
                       sec.first_page_footer, sec.even_page_header,
                       sec.even_page_footer):
                try:
                    for p in hf.paragraphs:
                        if p.style is not None:
                            names.add(p.style.name)
                except Exception:                   # noqa: BLE001
                    pass
        self._used_styles = names
        return names

    @property
    def fonts_used(self) -> dict[str, set]:
        """{字体名: {出现位置标签}}"""
        if self._fonts_used is not None:
            return self._fonts_used
        used: dict[str, set] = {}

        def note(name, where):
            if name and name.strip():
                used.setdefault(name.strip(), set()).add(where)

        style_names = self.used_style_names
        doc = self.doc

        # 样式层（仅被引用的样式）
        for style in doc.styles:
            if style.name not in style_names:
                continue
            el = style.element
            rPr = el.find(qn("w:rPr"))
            if rPr is None:
                continue
            rf = rPr.find(qn("w:rFonts"))
            if rf is None:
                continue
            for attr in ("w:ascii", "w:eastAsia", "w:hAnsi"):
                note(rf.get(qn(attr)), "样式 %s" % style.name)

        # 正文与表格的 run 层
        for p in self.doc.paragraphs:
            for r in p.runs:
                rPr = r._element.rPr
                if rPr is None:
                    continue
                rf = rPr.find(qn("w:rFonts"))
                if rf is None:
                    continue
                for attr in ("w:ascii", "w:eastAsia", "w:hAnsi"):
                    note(rf.get(qn(attr)), "正文")
        for t in self.doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            rPr = r._element.rPr
                            if rPr is None:
                                continue
                            rf = rPr.find(qn("w:rFonts"))
                            if rf is None:
                                continue
                            for attr in ("w:ascii", "w:eastAsia", "w:hAnsi"):
                                note(rf.get(qn(attr)), "表格")

        self._fonts_used = used
        return used

    def bad_crc_entries(self) -> list[str]:
        """逐条读一遍，暴露损坏的成员（图片/嵌入对象损坏的硬证据）"""
        if self._bad_crc is not None:
            return self._bad_crc
        bad = []
        for n in self.names:
            if n.endswith("/"):
                continue
            try:
                with self.zf.open(n) as f:
                    while f.read(65536):
                        pass
            except Exception:
                bad.append(n)
        self._bad_crc = bad
        return bad

    def close(self):
        try:
            self.zf.close()
        except Exception:
            pass


# ------------------------------------------------------------------ 报告
class Report:
    def __init__(self, path: str):
        self.path = path
        self.findings: list[Finding] = []
        self.rules_run: list[str] = []
        self.rules_skipped: list[str] = []
        self.error = None          # 无法解析时的原因

    def add(self, items: Iterable[Finding]):
        self.findings.extend(items)

    def counts(self):
        c = {"error": 0, "warning": 0, "info": 0}
        for f in self.findings:
            c[f.severity] = c.get(f.severity, 0) + 1
        return c

    def exit_code(self):
        if self.error:
            return 2
        return 1 if self.counts()["error"] else 0

    def to_dict(self):
        return {
            "tool": "docx-inspector",
            "file": self.path,
            "error": self.error,
            "counts": self.counts(),
            "rules_run": self.rules_run,
            "rules_skipped": self.rules_skipped,
            "findings": [f.to_dict() for f in self.findings],
        }


def run_rules(ctx: Ctx, report: Report, include: set[str], exclude: set[str]):
    """按注册顺序执行规则，单条规则异常不影响整体"""
    for rid, meta in RULES.items():
        if include and rid not in include:
            report.rules_skipped.append(rid)
            continue
        if rid in exclude:
            report.rules_skipped.append(rid)
            continue
        report.rules_run.append(rid)
        try:
            out = meta.fn(ctx) or []
        except Exception as e:                       # noqa: BLE001
            out = [Finding(rid, "warning",
                           "规则执行异常：%s: %s" % (type(e).__name__, e),
                           fixable=False)]
        report.add(out)
