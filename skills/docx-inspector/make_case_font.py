# -*- coding: utf-8 -*-
"""第三个用例：引用本机不存在的字体（仿宋_GB2312），验证 C10"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(3.17)


def set_font(run, ea, size=12):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:eastAsia"), ea)
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")


p = d.add_paragraph()
set_font(p.add_run("这段用仿宋_GB2312，本机未安装该字体。"), "仿宋_GB2312", 12)

p2 = d.add_paragraph()
set_font(p2.add_run("这段用楷体，本机已安装。"), "楷体", 12)

p3 = d.add_paragraph()
set_font(p3.add_run("这段又用仿宋_GB2312，制造正文字体不一致。"), "仿宋_GB2312", 12)

out = os.path.join(BASE, "case_font.docx")
d.save(out)
print("saved:", out, os.path.getsize(out), "bytes")
