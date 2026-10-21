# -*- coding: utf-8 -*-
"""生成 auto 布局的超宽表格 fixture（Word 会压缩，用于验证 B7 降级为 warning）"""
import os
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
WIDTHS = [5.0, 5.0, 5.0, 5.0]          # 20cm vs 版心 14.66cm

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(3.17)
sec.top_margin = sec.bottom_margin = Cm(2.54)

d.add_paragraph("auto 布局超宽表：声明 20cm，Word 渲染时会压缩回版心")

t = d.add_table(rows=3, cols=4)
t.style = "Table Grid"
# 注意：不要设 autofit=False —— python-docx 会写入 <w:tblLayout w:type="fixed"/>

# 重建 tblGrid 为真实 20cm，但不写 tblLayout（保持 auto）
tbl = t._tbl
tblPr = tbl.tblPr
old = tbl.find(qn("w:tblGrid"))
if old is not None:
    tbl.remove(old)
grid = OxmlElement("w:tblGrid")
for w in WIDTHS:
    gc = OxmlElement("w:gridCol")
    gc.set(qn("w:w"), str(int(w * 567)))
    grid.append(gc)
tblPr.addnext(grid)

for row in t.rows:
    for i, cell in enumerate(row.cells):
        cell.width = Cm(WIDTHS[i])
        tcPr = cell._tc.get_or_add_tcPr()
        for ow in tcPr.findall(qn("w:tcW")):
            tcPr.remove(ow)
        tcw = OxmlElement("w:tcW")
        tcw.set(qn("w:w"), str(int(WIDTHS[i] * 567)))
        tcw.set(qn("w:type"), "dxa")
        tcPr.insert(0, tcw)
        cell.text = "第%d列内容占位" % (i + 1)

out = os.path.join(BASE, "case_auto_overflow.docx")
d.save(out)
print("saved:", out, os.path.getsize(out), "bytes")
