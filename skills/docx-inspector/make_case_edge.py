# -*- coding: utf-8 -*-
"""为两个第三方发现的缺陷造 fixture：
  case_no_pgmar.docx       — sectPr 缺 pgMar（此前会导致整篇 exit 2）
  case_merged_header.docx  — 首行横向合并单元格（此前 D15 误报）
"""
import os

from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))


def set_grid(table, widths_cm):
    tbl = table._tbl
    tblPr = tbl.tblPr
    old = tbl.find(qn("w:tblGrid"))
    if old is not None:
        tbl.remove(old)
    grid = OxmlElement("w:tblGrid")
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w * 567)))
        grid.append(gc)
    tblPr.addnext(grid)
    for row in table.rows:
        for ci, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            for ow in tcPr.findall(qn("w:tcW")):
                tcPr.remove(ow)
            tcw = OxmlElement("w:tcW")
            tcw.set(qn("w:w"), str(int(widths_cm[min(ci, len(widths_cm) - 1)] * 567)))
            tcw.set(qn("w:type"), "dxa")
            tcPr.insert(0, tcw)


# ---------------------------------------------------------------- fixture 1
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(3.17)
sec.top_margin = sec.bottom_margin = Cm(2.54)

d.add_paragraph("页面属性缺失样例：sectPr 只有 pgSz，没有 pgMar")

t = d.add_table(rows=3, cols=3)
t.style = "Table Grid"
set_grid(t, [4.0, 5.0, 5.0])
for r in t.rows:
    for c in r.cells:
        c.text = "内容"

# 关键：删掉 pgMar，模拟真实文档的写法
sectPr = sec._sectPr
pgMar = sectPr.find(qn("w:pgMar"))
if pgMar is not None:
    sectPr.remove(pgMar)
    print("已删除 pgMar")

out1 = os.path.join(BASE, "case_no_pgmar.docx")
d.save(out1)
print("saved:", out1, os.path.getsize(out1), "bytes")

# ---------------------------------------------------------------- fixture 2
d2 = Document()
sec2 = d2.sections[0]
sec2.page_width, sec2.page_height = Cm(21.0), Cm(29.7)
sec2.left_margin = sec2.right_margin = Cm(3.17)

d2.add_paragraph("首行横向合并单元格样例：不应触发 D15")

t2 = d2.add_table(rows=3, cols=4)
t2.style = "Table Grid"
# 首行 4 列合并成 1 个单元格（表头跨列）
t2.cell(0, 0).merge(t2.cell(0, 3))
t2.cell(0, 0).text = "合并表头（跨 4 列）"
for ci in range(4):
    t2.cell(1, ci).text = "B%d" % (ci + 1)
    t2.cell(2, ci).text = "C%d" % (ci + 1)

# tblGrid 与各 tcW 保持一致：合并单元格的 tcW = 整表宽
W = [3.4, 3.4, 3.4, 3.4]
tbl2 = t2._tbl
tblPr2 = tbl2.tblPr
old2 = tbl2.find(qn("w:tblGrid"))
if old2 is not None:
    tbl2.remove(old2)
grid2 = OxmlElement("w:tblGrid")
for w in W:
    gc = OxmlElement("w:gridCol")
    gc.set(qn("w:w"), str(int(w * 567)))
    grid2.append(gc)
tblPr2.addnext(grid2)

# 首行：单个 tc，宽 = 整表；其余行：4 个 tc
total_tw = int(sum(W) * 567)
for ri, row in enumerate(t2.rows):
    tcs = row._tr.findall(qn("w:tc"))
    for tc in tcs:
        tcPr = tc.find(qn("w:tcPr"))
        if tcPr is None:
            tcPr = OxmlElement("w:tcPr")
            tc.insert(0, tcPr)
        for ow in tcPr.findall(qn("w:tcW")):
            tcPr.remove(ow)
        tcw = OxmlElement("w:tcW")
        if ri == 0:
            tcw.set(qn("w:w"), str(total_tw))       # 合并单元格 = 整表宽
        else:
            tcw.set(qn("w:w"), str(int(W[0] * 567)))
        tcw.set(qn("w:type"), "dxa")
        tcPr.insert(0, tcw)

out2 = os.path.join(BASE, "case_merged_header.docx")
d2.save(out2)
print("saved:", out2, os.path.getsize(out2), "bytes")
