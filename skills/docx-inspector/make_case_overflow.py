# -*- coding: utf-8 -*-
"""造一个"真溢出"样本：tblLayout=fixed + 明确 tcW，逼 Word 不压缩"""
import os
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(3.17)
sec.top_margin = sec.bottom_margin = Cm(2.54)

# 可用宽 14.66cm，这里做 4 列 × 5cm = 20cm，明显超
WIDTHS = [5.0, 5.0, 5.0, 5.0]
d.add_paragraph("溢出测试：表格总宽 20cm vs 版心 14.66cm")

t = d.add_table(rows=4, cols=4)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.LEFT
t.autofit = False

tbl = t._tbl
tblPr = tbl.tblPr

# fixed 布局
layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
tblPr.append(layout)
# tblW
tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(WIDTHS) * 567))); tw.set(qn("w:type"), "dxa")
tblPr.append(tw)
# 重建 tblGrid
old = tbl.find(qn("w:tblGrid"))
if old is not None:
    tbl.remove(old)
grid = OxmlElement("w:tblGrid")
for w in WIDTHS:
    gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(int(w * 567)))
    grid.append(gc)
tblPr.addnext(grid)

for row in t.rows:
    for i, cell in enumerate(row.cells):
        cell.width = Cm(WIDTHS[i])
        tcPr = cell._tc.get_or_add_tcPr()
        for ow in tcPr.findall(qn("w:tcW")):
            tcPr.remove(ow)
        tcw = OxmlElement("w:tcW"); tcw.set(qn("w:w"), str(int(WIDTHS[i] * 567))); tcw.set(qn("w:type"), "dxa")
        tcPr.insert(0, tcw)
        cell.text = "第%d列超宽内容占位文字" % (i + 1)

out = os.path.join(BASE, "case_overflow.docx")
d.save(out)
print("saved:", out, os.path.getsize(out), "bytes")
