# -*- coding: utf-8 -*-
"""生成完全脱敏的合规 fixture（case_clean.docx）

要求：
  * 不含任何真实公司 / 项目 / 人名
  * 通过全部默认规则（无 B7 / C10 / D15 / E20）
  * 保留一定结构复杂度（标题、表格、列表），使测试有意义
"""
import os

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
WIDTHS = [2.0, 4.0, 4.0, 4.0]          # 合计 14.0cm < 版心 14.66cm

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(3.17)
sec.top_margin = sec.bottom_margin = Cm(2.54)

# 不显式设置任何字体 —— 让 C10 无话可说（也不制造字体缺失）
p = d.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("示例技术评审报告")
r.bold = True
r.font.size = Pt(16)

d.add_paragraph()
d.add_paragraph("一、概述").runs[0].bold = True
d.add_paragraph(
    "本文档为示例文档，仅用于验证排版检查工具的行为，不包含任何真实业务信息。"
    "文中单位、系统、人名均为虚构占位。")

d.add_paragraph("二、检查项").runs[0].bold = True

t = d.add_table(rows=5, cols=4)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER

tbl = t._tbl
tblPr = tbl.tblPr
old = tbl.find(qn("w:tblGrid"))
if old is not None:
    tbl.remove(old)
grid = OxmlElement("w:tblGrid")
for w in WIDTHS:
    gc = OxmlElement("w:gridCol")
    gc.set(qn("w:w"), str(int(w * 567)))
    grid.append(gc)
tblPr.addnext(grid)

header = ["序号", "检查项", "结果", "备注"]
rows = [
    header,
    ["1", "外观检查", "通过", "无异常"],
    ["2", "功能检查", "通过", "无异常"],
    ["3", "性能检查", "通过", "无异常"],
    ["4", "安全检查", "通过", "无异常"],
]
for ri, row in enumerate(rows):
    for ci, text in enumerate(row):
        cell = t.cell(ri, ci)
        cell.text = text
        tcPr = cell._tc.get_or_add_tcPr()
        for ow in tcPr.findall(qn("w:tcW")):
            tcPr.remove(ow)
        tcw = OxmlElement("w:tcW")
        tcw.set(qn("w:w"), str(int(WIDTHS[ci] * 567)))
        tcw.set(qn("w:type"), "dxa")
        tcPr.insert(0, tcw)

d.add_paragraph()
d.add_paragraph("三、结论").runs[0].bold = True
d.add_paragraph("综上所述，示例系统各项指标均符合预期，建议通过评审。")

d.add_paragraph("后续事项：").runs[0].bold = True
for i, item in enumerate(["归档评审记录", "同步至项目看板", "安排复审"], 1):
    p = d.add_paragraph(style="List Paragraph")
    p.add_run(item)
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId"); numId.set(qn("w:val"), "1")
    numPr.append(ilvl); numPr.append(numId)
    pPr.insert_element_before(
        numPr, "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs",
        "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
        "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
        "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")

out = os.path.join(BASE, "case_clean.docx")
d.save(out)
print("saved:", out, os.path.getsize(out), "bytes")
